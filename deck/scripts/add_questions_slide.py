#!/usr/bin/env python3
"""Insert a '3-layer natural-language questions' slide after the demo slide,
matching the deck's Searce style, then renumber all footer page numbers.
Operates in place on deck/TSUK-Logistics-Cost-Analytics.pptx.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

HERE = os.path.dirname(__file__)
DECK = os.path.join(HERE, "..", "TSUK-Logistics-Cost-Analytics.pptx")
IMG = os.path.join(HERE, "..", "images")

BLUE = RGBColor(0x00, 0x64, 0xFF); BLUEDK = RGBColor(0x00, 0x51, 0xCB)
NAVY = RGBColor(0x00, 0x26, 0x59); BODY = RGBColor(0x47, 0x55, 0x69)
SUB = RGBColor(0x5F, 0x63, 0x68); MUTED = RGBColor(0x9A, 0xA0, 0xA6)
CARDLN = RGBColor(0xE0, 0xE6, 0xF6); PANEL = RGBColor(0xF4, 0xF7, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF); H = "Poppins"

prs = Presentation(DECK)
INSERT_AFTER = 6  # after slide 6 (the demo) -> new slide becomes position 7


def rect(s, l, t, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, rad=0.06):
    sp = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0] = rad
        except Exception: pass
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(lw)
    return sp


def text(s, l, t, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", align)
        if para.get("after") is not None: p.space_after = Pt(para["after"])
        if para.get("lh") is not None: p.line_spacing = para["lh"]
        for run in para["runs"]:
            r = p.add_run(); r.text = run[0]
            r.font.name = H; r.font.size = Pt(run[1]); r.font.bold = run[3]
            r.font.italic = run[4] if len(run) > 4 else False
            r.font.color.rgb = run[2]
    return tb


def pic(s, name, l, t, w):
    path = os.path.join(IMG, name); iw, ih = Image.open(path).size
    s.shapes.add_picture(path, Inches(l), Inches(t), Inches(w), Inches(w * ih / iw))


# ── build the new slide (blank, reusing demo slide's layout) ──
layout = prs.slides[INSERT_AFTER - 1].slide_layout
ns = prs.slides.add_slide(layout)
for ph in list(ns.placeholders):
    ph._element.getparent().remove(ph._element)
ns.background.fill.solid(); ns.background.fill.fore_color.rgb = WHITE

rect(ns, 0.34, 0.42, 0.022, 4.58, fill=BLUE)                       # hairline
text(ns, 0.6, 0.42, 8.8, 0.3, [{"runs": [("The demo · ask anything", 12.5, BLUE, True)]}])
text(ns, 0.6, 0.66, 8.9, 0.7, [{"runs": [("Ask in plain English — then drill deeper", 26, NAVY, True)], "lh": 1.0}])
text(ns, 0.6, 1.34, 8.9, 0.4, [{"runs": [
    ("Three questions, one flowing conversation — every answer grounded in the model, ", 13, BODY, False),
    ("never invented.", 13, NAVY, True)]}])

cards = [
    ("01", "START BROAD", "“What’s our total logistics cost, split by mode and provider?”",
     "£81.4M across rail, UK road and EU road"),
    ("02", "GO DEEPER", "“How does cost per tonne compare across rail, UK road and EU road?”",
     "EU road £55.31 vs rail £6.57 — 8.4× more per tonne"),
    ("03", "MAKE IT ACTIONABLE", "“Which lanes and carriers drive the most of that cost?”",
     "Ranked targets to consolidate or renegotiate"),
]
x, cw, ch, gap, y0 = 0.6, 9.0, 0.98, 0.17, 1.84
for i, (num, tag, q, ret) in enumerate(cards):
    yy = y0 + i * (ch + gap)
    rect(ns, x, yy, cw, ch, fill=PANEL, line=CARDLN, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rad=0.1)
    rect(ns, x, yy, 0.06, ch, fill=BLUE)
    text(ns, x + 0.24, yy, 0.85, ch, [{"runs": [(num, 28, BLUE, True)]}], anchor=MSO_ANCHOR.MIDDLE)
    text(ns, x + 1.15, yy + 0.13, 7.6, 0.25, [{"runs": [(tag, 10, BLUE, True)]}])
    text(ns, x + 1.15, yy + 0.35, 7.7, 0.32, [{"runs": [(q, 15.5, NAVY, True)]}])
    text(ns, x + 1.15, yy + 0.67, 7.7, 0.28, [{"runs": [
        ("Returns  ", 11.5, BLUE, True), (ret, 12.5, BODY, False)]}])

# footer (logo + a number that the renumber pass will correct)
pic(ns, "searce-dark.png", 0.6, 5.18, 0.82)
text(ns, 8.3, 5.26, 1.2, 0.25, [{"runs": [("07", 9, MUTED, False)]}], align=PP_ALIGN.RIGHT)

# ── move new slide to position 7 (index 6) ──
lst = prs.slides._sldIdLst
ids = list(lst); new_id = ids[-1]
lst.remove(new_id); lst.insert(INSERT_AFTER, new_id)

# ── renumber every footer page-number to its final position ──
for i, s in enumerate(prs.slides):
    for sh in s.shapes:
        if sh.has_text_frame and Emu(sh.top).inches > 5.0 and Emu(sh.left).inches > 7.5:
            tf = sh.text_frame
            if tf.text.strip().isdigit():
                tf.paragraphs[0].runs[0].text = f"{i + 1:02d}"

prs.save(DECK)
print("inserted at pos", INSERT_AFTER + 1, "| total slides:", len(prs.slides._sldIdLst))
