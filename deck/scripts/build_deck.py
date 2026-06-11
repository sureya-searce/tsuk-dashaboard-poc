#!/usr/bin/env python3
"""Build the Searce-branded TSUK Logistics Cost Analytics deck (10 slides, 16:9)
in the REAL Searce system (from the BQ-Graph POV deck): editorial / minimal,
white slides, blue hairline motif, blue eyebrows, navy Poppins headlines,
white cards with big blue numbers, light-blue callouts, logo bottom-left.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

HERE = os.path.dirname(__file__)
IMG = os.path.join(HERE, "..", "images")
OUT = os.path.join(HERE, "..", "TSUK-Logistics-Cost-Analytics.pptx")

# ── Searce palette ──
BLUE   = RGBColor(0x00, 0x64, 0xFF)
BLUEDK = RGBColor(0x00, 0x51, 0xCB)
NAVY   = RGBColor(0x00, 0x26, 0x59)
INK    = RGBColor(0x20, 0x21, 0x24)
BODY   = RGBColor(0x47, 0x55, 0x69)
SUB    = RGBColor(0x5F, 0x63, 0x68)
MUTED  = RGBColor(0x9A, 0xA0, 0xA6)
CARDLN = RGBColor(0xE0, 0xE6, 0xF6)
CALL   = RGBColor(0xE6, 0xF0, 0xFF)
PANEL  = RGBColor(0xF4, 0xF7, 0xFB)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
H = "Poppins"          # Google Slides has Poppins natively

prs = Presentation()
prs.slide_width = Inches(10); prs.slide_height = Inches(5.625)
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE
    return s


def rect(s, l, t, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, rad=0.06):
    sp = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0] = rad
        except Exception: pass
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = line; sp.line.width = Pt(lw)
    return sp


def text(s, l, t, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", align)
        if para.get("after") is not None: p.space_after = Pt(para["after"])
        if para.get("before") is not None: p.space_before = Pt(para["before"])
        if para.get("lh") is not None: p.line_spacing = para["lh"]
        for run in para["runs"]:
            txt, sz, col, bold = run[0], run[1], run[2], run[3]
            ital = run[4] if len(run) > 4 else False
            r = p.add_run(); r.text = txt
            r.font.name = H; r.font.size = Pt(sz); r.font.bold = bold
            r.font.italic = ital; r.font.color.rgb = col
    return tb


def pic(s, name, l, t, w=None, h=None):
    path = os.path.join(IMG, name); iw, ih = Image.open(path).size
    if w and not h: h = w * ih / iw
    if h and not w: w = h * iw / ih
    return s.shapes.add_picture(path, Inches(l), Inches(t),
                                Inches(w) if w else None, Inches(h) if h else None)


def hairline(s, top=0.42, bottom=5.0):
    rect(s, 0.34, top, 0.022, bottom - top, fill=BLUE)


def head(s, eyebrow, title, sub=None):
    hairline(s)
    text(s, 0.6, 0.42, 8.8, 0.3, [{"runs": [(eyebrow, 12.5, BLUE, True)]}])
    text(s, 0.6, 0.66, 8.9, 0.7, [{"runs": [(title, 26, NAVY, True)], "lh": 1.0}])
    if sub:
        text(s, 0.6, 1.34, 8.8, 0.4, [{"runs": sub}])


def footer(s, n):
    pic(s, "searce-dark.png", 0.6, 5.18, w=0.82)
    text(s, 8.3, 5.26, 1.2, 0.25, [{"runs": [(f"{n:02d}", 9, MUTED, False)]}], align=PP_ALIGN.RIGHT)


# ════════════════════════════════ 1 · TITLE
s = slide()
hairline(s, 0.42, 5.2)
text(s, 0.62, 0.5, 3, 0.7, [
    {"runs": [("SEARCE  ×  TATA STEEL UK", 13, BLUE, True)], "after": 2},
    {"runs": [("FY26  ·  Logistics cost transformation", 13, BLUE, False)]}])
text(s, 0.62, 2.55, 6.6, 2.2, [
    {"runs": [("From three spreadsheets", 40, NAVY, True)], "lh": 1.04},
    {"runs": [("to one conversation.", 40, NAVY, True)], "lh": 1.04}])
pic(s, "searce-dark.png", 0.62, 5.0, w=1.0)
pic(s, "tatasteel.png", 7.9, 0.55, w=1.5)

# ════════════════════════════════ 2 · THE PROBLEM
s = slide()
head(s, "The problem", "Logistics cost you can’t see",
     [("Cost lives in ", 13, BODY, False), ("three disconnected spreadsheets", 13, NAVY, True),
      (", emailed monthly — each a different shape.", 13, BODY, False)])
feeds = [("Rail — DB Cargo", "1 row / movement", "Costs split across haulage, fuel and cancellation"),
         ("Road UK", "1 row / leg", "Multi-leg orders; purchase vs sales cost"),
         ("Road EU — imports", "1 row / charge line", "~4 lines per shipment; customer code ≠ shipment")]
cx = [0.6, 3.78, 6.96]; cw = 2.66
for x, (nm, grain, why) in zip(cx, feeds):
    rect(s, x, 2.0, cw, 2.05, fill=WHITE, line=CARDLN, lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rad=0.05)
    text(s, x + 0.22, 2.2, cw - 0.44, 0.4, [{"runs": [(nm, 14.5, NAVY, True)]}])
    rect(s, x + 0.22, 2.62, 0.45, 0.03, fill=BLUE)
    text(s, x + 0.22, 2.8, cw - 0.44, 1.1, [
        {"runs": [(grain, 13, BLUE, True)], "after": 7},
        {"runs": [(why, 12.5, BODY, False)], "lh": 1.1}])
rect(s, 0.6, 4.35, 9.0, 0.66, fill=CALL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rad=0.1)
text(s, 0.85, 4.35, 8.5, 0.66, [{"runs": [
    ("Nobody can answer, across all three at once:  ", 13.5, NAVY, True),
    ("what’s our cost per tonne, where is it leaking, are we paying to move air?", 13.5, BODY, False, True)]}],
    anchor=MSO_ANCHOR.MIDDLE)
footer(s, 2)

# ════════════════════════════════ 3 · BEFORE / AFTER
s = slide()
head(s, "Before vs after", "From 6 days to minutes")
pic(s, "before-after.png", 1.25, 1.5, w=7.5)
footer(s, 3)

# ════════════════════════════════ 4 · CONCEPTUAL
s = slide()
head(s, "How it works", "One file in, answers out")
pic(s, "arch-conceptual.png", 0.5, 1.7, w=9.0)
footer(s, 4)

# ════════════════════════════════ 5 · TECHNICAL
s = slide()
head(s, "Architecture", "Event-driven on Google Cloud")
pic(s, "arch-technical.png", 0.62, 1.5, w=8.76)
footer(s, 5)

# ════════════════════════════════ 6 · DEMO STEPS
s = slide()
head(s, "The demo", "Watch it happen, live")
steps = [("01", "Start blank", "Truncate the data — the dashboard shows nothing"),
         ("02", "Drop a file", "Land a supplier file in the cloud landing zone"),
         ("03", "It runs itself", "Ingest, unify, KPIs, anomalies — no human"),
         ("04", "Ask in English", "Conversational Analytics answers from the model")]
y0 = 1.7; rh = 0.74
for i, (n, hd, sub) in enumerate(steps):
    yy = y0 + i * (rh + 0.12)
    rect(s, 0.6, yy, 9.0, rh, fill=PANEL, line=CARDLN, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rad=0.12)
    text(s, 0.85, yy, 0.9, rh, [{"runs": [(n, 26, BLUE, True)]}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.85, yy, 7.5, rh, [{"runs": [
        (hd + "    ", 15.5, NAVY, True), (sub, 14, BODY, False)]}], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 6)

# ════════════════════════════════ 7 · COST STORY
s = slide()
head(s, "The demo · cost story", "The insight that was hidden")
pic(s, "cost-insight.png", 0.45, 1.55, w=5.15)
text(s, 5.95, 1.7, 3.6, 3.3, [
    {"runs": [("European road freight costs", 18, NAVY, True)], "lh": 1.05},
    {"runs": [("8.4× per tonne", 32, BLUE, True)], "before": 6, "after": 2},
    {"runs": [("what rail does.", 18, NAVY, True)], "after": 12, "lh": 1.05},
    {"runs": [("Invisible while the three feeds lived apart — the most actionable cost fact on the table.", 12.5, BODY, False)], "lh": 1.15},
], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 7)

# ════════════════════════════════ 8 · THE RECEIPTS
s = slide()
head(s, "The demo · trust", "Every number ties to source")
stats = [("£81.4M", "logistics spend analysed"), ("142,090", "movements unified"), ("£0", "reconciliation gap")]
sx = [0.6, 3.78, 6.96]; sw = 2.66
for x, (big, lab) in zip(sx, stats):
    rect(s, x, 1.6, sw, 1.3, fill=PANEL, line=CARDLN, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rad=0.07)
    text(s, x, 1.78, sw, 0.7, [{"runs": [(big, 38, BLUE, True)]}], align=PP_ALIGN.CENTER)
    text(s, x, 2.5, sw, 0.35, [{"runs": [(lab, 12.5, BODY, False)]}], align=PP_ALIGN.CENTER)
recon = [("Rail (DB Cargo)", "£27,124,697"), ("Road UK", "£36,367,196"), ("Road EU (imports)", "£17,930,073")]
ry = 3.4
text(s, 0.6, ry - 0.3, 9, 0.3, [{"runs": [("Computed total  =  source control total, per feed", 13, NAVY, True)]}])
for i, (nm, val) in enumerate(recon):
    yy = ry + i * 0.46
    rect(s, 0.6, yy, 9.0, 0.42, fill=PANEL if i % 2 == 0 else WHITE)
    text(s, 0.8, yy, 4, 0.42, [{"runs": [(nm, 12.5, INK, False)]}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 4.7, yy, 2.4, 0.42, [{"runs": [(val, 12.5, BODY, False)]}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 7.2, yy, 2.3, 0.42, [{"runs": [("✓  matches to the penny", 12.5, BLUE, True)]}], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 8)

# ════════════════════════════════ 9 · CLOSE / NEXT STEPS
s = slide()
hairline(s, 0.42, 5.2)
text(s, 0.6, 0.5, 4, 0.3, [{"runs": [("The close", 12.5, BLUE, True)]}])
text(s, 0.6, 1.0, 8.8, 1.5, [
    {"runs": [("One file drop. Answers in plain English.", 32, NAVY, True)], "lh": 1.05}])
text(s, 0.62, 2.35, 8.8, 0.5, [{"runs": [("The control panel for the cost you can actually move.", 16, BLUE, False)]}])
nexts = [("Productionise", "IaC, scheduled supplier feeds, CI/CD"),
         ("Broaden coverage", "onboard more carriers & modes"),
         ("Source-data quality", "quarantine + alerting on gaps")]
nx = [0.6, 3.78, 6.96]; nw = 2.66
text(s, 0.6, 3.25, 4, 0.3, [{"runs": [("NEXT STEP", 11, BLUE, True)]}])
for x, (hd, sub) in zip(nx, nexts):
    rect(s, x, 3.6, nw, 1.05, fill=WHITE, line=CARDLN, lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE, rad=0.07)
    rect(s, x, 3.6, 0.06, 1.05, fill=BLUE)
    text(s, x + 0.24, 3.76, nw - 0.42, 0.9, [
        {"runs": [(hd, 14.5, NAVY, True)], "after": 4},
        {"runs": [(sub, 11.5, BODY, False)], "lh": 1.1}])
footer(s, 9)

# ════════════════════════════════ 10 · QUESTIONS
s = slide()
hairline(s, 0.42, 5.2)
text(s, 0.6, 2.0, 9, 1.2, [{"runs": [("Questions?", 50, NAVY, True)]}])
text(s, 0.62, 3.25, 9, 0.5, [{"runs": [("Thank you  ·  ", 18, BLUE, True), ("Searce × Tata Steel UK", 18, BODY, False)]}])
pic(s, "searce-dark.png", 0.6, 4.6, w=1.1)

prs.save(OUT)
print("saved", os.path.abspath(OUT), "slides:", len(prs.slides._sldIdLst))
